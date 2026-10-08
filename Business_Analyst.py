"""
NexaMart / Generic Business Intelligence + Financial RAG Assistant
---------------------------------------------------------------
Streamlit application supporting:
- PDF, TXT, CSV, Excel, Word and common chart/image files
- RAG retrieval with configurable chunk size, overlap and Top-K
- Groq LLM selection and API-key input
- Numerical CSV analysis using pandas
- Financial/business questions across uploaded documents
- Source/page/file citations
- Year comparison, revenue, expenses, product/customer/inventory analysis
- Optional chart/image OCR when Tesseract is installed

Install:
    pip install -r requirements.txt

Run:
    streamlit run app.py
"""

import os
import re
import io
import json
import math
import hashlib
from pathlib import Path
from typing import List, Dict, Any, Tuple

import pandas as pd
import numpy as np
import streamlit as st

# ChromaDB persistent vector database
try:
    import chromadb
except ImportError:
    chromadb = None

# Excel / Word support
try:
    from docx import Document
except ImportError:
    Document = None

# PDF
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

# Embeddings / vector retrieval
try:
    from sentence_transformers import SentenceTransformer
except ImportError:
    SentenceTransformer = None

# Groq
try:
    from groq import Groq
except ImportError:
    Groq = None

# Optional image OCR
try:
    from PIL import Image
except ImportError:
    Image = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

import os
from dotenv import load_dotenv

# Load Gorq Key from .env file as variable.

load_dotenv()
# Get Groq API key
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
#print("API key loaded:", GROQ_API_KEY is not None)


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Business Intelligence & Financial RAG Assistant",
    page_icon="📊",
    layout="wide",
)

st.markdown(
    """
    <style>
    .main {
        background-color: #f7fbff;
    }
    .stChatMessage {
        border-radius: 12px;
    }
    .source-box {
        padding: 10px;
        border-radius: 8px;
        background: #eef6ff;
        margin-top: 5px;
        margin-bottom: 5px;
    }

    /* =========================
       ALL BUTTONS
       ========================= */
    .stButton > button {
        background-color: #0284c7;
        color: white;
        border: none;
        border-radius: 10px;
        padding: 0.6rem 1.2rem;
        font-weight: 600;
    }

    .stButton > button:hover {
        background-color: #0369a1;
        color: white;
        border: none;
    }

    /* =========================
       PRIMARY BUTTON
       ========================= */
    .stButton > button[kind="primary"] {
        background-color: #16a34a;
        color: white;
        border-radius: 10px;
        font-weight: 700;
    }

    .stButton > button[kind="primary"]:hover {
        background-color: #15803d;
        color: white;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "documents": [],
    "chunks": [],
    "embeddings": None,  # kept for backward compatibility; Chroma is now the vector store
    "chroma_collection": None,
    "tables": {},
    "file_metadata": [],
    "messages": [],
    "indexed": False,
    "embedder": None,
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.title("⚙️ RAG Configuration")

groq_api_input = st.sidebar.text_input(
    "Groq API Key",
    type="password",
    help = "Your Groq API key is used only for the current Streamlit session.",
 # We write following value from online " https://console.groq.com/keys" on this site and create API key.   
   
    #value = "directly paste value here for testing but it will apear unsecure on frontend textfield"

   )

model_options = [
   
    "openai/gpt-oss-120b",
    "openai/gpt-oss-20b",
]

selected_model = st.sidebar.selectbox(
    "Select LLM Model",
    model_options,
)

chunk_size = st.sidebar.slider(
    "Chunk Size",
    min_value=300,
    max_value=5000,
    value=1200,
    step=100,
)

chunk_overlap = st.sidebar.slider(
    "Chunk Overlap",
    min_value=0,
    max_value=1000,
    value=200,
    step=50,
)

top_k = st.sidebar.slider(
    "Top-K Retrieved Chunks",
    min_value=1,
    max_value=15,
    value=5,
    step=1,
)

temperature = st.sidebar.slider(
    "LLM Temperature",
    min_value=0.0,
    max_value=1.0,
    value=0.1,
    step=0.05,
)

st.sidebar.markdown("---")

st.sidebar.info(
    "Tip: For financial/business questions, upload the CSV files as well as "
    "annual reports, policies, Excel and Word files. CSV/Excel analysis is calculated with pandas; "
    "the LLM is used to explain the results. "
    "Document chunks are stored persistently in ChromaDB for semantic retrieval."
)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def normalize_name(name: str) -> str:
    """Normalize a column name for easier automatic detection."""
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Basic cleanup without changing business meaning."""
    df = df.copy()
    df.columns = [normalize_name(c) for c in df.columns]

    for col in df.columns:
        if df[col].dtype == "object":
            df[col] = df[col].astype(str).str.strip()

    return df


def detect_date_columns(df: pd.DataFrame) -> List[str]:
    candidates = []
    for col in df.columns:
        name = normalize_name(col)
        if any(x in name for x in [
            "date", "year", "month", "period", "time"
        ]):
            candidates.append(col)

    # Also inspect object columns for parseable dates.
    for col in df.select_dtypes(include=["object"]).columns:
        if col not in candidates:
            sample = df[col].dropna().head(30)
            if len(sample):
                parsed = pd.to_datetime(sample, errors="coerce")
                if parsed.notna().mean() >= 0.7:
                    candidates.append(col)

    return list(dict.fromkeys(candidates))


def find_column(df: pd.DataFrame, patterns: List[str]) -> str | None:
    """Find the most likely column matching business concepts."""
    normalized = {c: normalize_name(c) for c in df.columns}

    # Exact / strong match first.
    for p in patterns:
        p = normalize_name(p)
        for col, ncol in normalized.items():
            if ncol == p:
                return col

    for p in patterns:
        p = normalize_name(p)
        for col, ncol in normalized.items():
            if p in ncol:
                return col

    return None


def parse_year_from_text(text: str) -> int | None:
    years = re.findall(r"\b(20\d{2})\b", text)
    if not years:
        return None
    return int(years[-1])


def detect_requested_years(question: str) -> List[int]:
    years = [int(x) for x in re.findall(r"\b(20\d{2})\b", question)]
    return sorted(set(years))


def infer_period(question: str) -> str:
    q = question.lower()

    if "this year" in q or "current year" in q:
        return "current"
    if "last year" in q or "previous year" in q or "prior year" in q:
        return "previous"
    return "explicit"


def numeric_columns(df: pd.DataFrame) -> List[str]:
    return list(df.select_dtypes(include=np.number).columns)


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_pdf(uploaded_file) -> Tuple[str, List[Dict[str, Any]]]:
    if PdfReader is None:
        raise RuntimeError("pypdf is not installed.")

    data = uploaded_file.getvalue()
    reader = PdfReader(io.BytesIO(data))

    pages = []
    full_text = []

    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(
                {
                    "text": text,
                    "source": uploaded_file.name,
                    "page": page_no,
                    "type": "pdf",
                }
            )
            full_text.append(text)

    return "\n\n".join(full_text), pages


def dataframe_to_text(df: pd.DataFrame, filename: str) -> str:
    """Create retrieval-friendly text from a CSV."""
    lines = [
        f"CSV FILE: {filename}",
        f"Rows: {len(df)}",
        f"Columns: {', '.join(df.columns)}",
        "",
        "COLUMN INFORMATION:",
    ]

    for col in df.columns:
        dtype = str(df[col].dtype)
        non_null = int(df[col].notna().sum())
        lines.append(f"- {col}: dtype={dtype}, non_null={non_null}")

    lines.append("")
    lines.append("SAMPLE DATA:")

    sample = df.head(100).to_string(index=False)
    lines.append(sample)

    return "\n".join(lines)


def extract_image_text(uploaded_file) -> Tuple[str, List[Dict[str, Any]]]:
    if Image is None:
        return "", []

    image = Image.open(io.BytesIO(uploaded_file.getvalue()))

    if pytesseract is None:
        return (
            "IMAGE/CHART FILE: "
            + uploaded_file.name
            + "\nOCR is not installed. Install pytesseract and Tesseract OCR "
              "to extract text from chart images.",
            [{
                "text": "Chart/image uploaded: " + uploaded_file.name,
                "source": uploaded_file.name,
                "page": None,
                "type": "image",
            }],
        )

    try:
        text = pytesseract.image_to_string(image)
    except Exception as exc:
        text = f"OCR failed: {exc}"

    return text, [{
        "text": text or f"Chart/image uploaded: {uploaded_file.name}",
        "source": uploaded_file.name,
        "page": None,
        "type": "image",
    }]


def extract_word_text(uploaded_file) -> Tuple[str, List[Dict[str, Any]]]:
    """Extract paragraphs and tables from a Word document."""
    if Document is None:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx")
    doc = Document(io.BytesIO(uploaded_file.getvalue()))
    parts = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table_no, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        parts.append(f"WORD TABLE {table_no}:\n" + "\n".join(rows))
    text = "\n\n".join(parts).strip()
    return text, [{"text": text or f"Word document uploaded: {uploaded_file.name}", "source": uploaded_file.name, "page": None, "type": "word"}]


def extract_excel_text(uploaded_file) -> Tuple[str, List[Dict[str, Any]], Dict[str, pd.DataFrame]]:
    """Extract all Excel worksheets as DataFrames and RAG text."""
    sheets = pd.read_excel(io.BytesIO(uploaded_file.getvalue()), sheet_name=None)
    documents, tables, all_text = [], {}, []
    for sheet_name, df in sheets.items():
        df = clean_dataframe(df)
        key = f"{uploaded_file.name} | {sheet_name}"
        tables[key] = df
        sheet_text = f"EXCEL SHEET: {sheet_name}\n{dataframe_to_text(df, key)}"
        all_text.append(sheet_text)
        documents.append({"text": sheet_text, "source": uploaded_file.name, "page": f"sheet: {sheet_name}", "type": "excel"})
    return "\n\n".join(all_text), documents, tables


def load_uploaded_files(uploaded_files):
    documents = []
    tables = {}
    metadata = []

    for file in uploaded_files:
        name = file.name
        suffix = Path(name).suffix.lower()

        try:
            if suffix == ".pdf":
                text, pages = extract_pdf(file)

                documents.extend(pages)

                metadata.append({
                    "file": name,
                    "type": "PDF",
                    "rows": None,
                    "status": "loaded",
                })

            elif suffix == ".csv":
                df = pd.read_csv(file)
                df = clean_dataframe(df)
                tables[name] = df

                text = dataframe_to_text(df, name)

                documents.append({
                    "text": text,
                    "source": name,
                    "page": None,
                    "type": "csv",
                })

                metadata.append({
                    "file": name,
                    "type": "CSV",
                    "rows": len(df),
                    "status": "loaded",
                })

            elif suffix in [".xlsx", ".xls"]:
                text, parts, excel_tables = extract_excel_text(file)
                documents.extend(parts)
                tables.update(excel_tables)
                metadata.append({"file": name, "type": "EXCEL", "rows": sum(len(df) for df in excel_tables.values()), "status": "loaded"})

            elif suffix in [".docx", ".doc"]:
                text, parts = extract_word_text(file)
                documents.extend(parts)
                metadata.append({"file": name, "type": "WORD", "rows": None, "status": "loaded"})

            elif suffix in [".txt", ".md"]:
                raw = file.getvalue().decode("utf-8", errors="ignore")

                documents.append({
                    "text": raw,
                    "source": name,
                    "page": None,
                    "type": "text",
                })

                metadata.append({
                    "file": name,
                    "type": "TXT/MD",
                    "rows": None,
                    "status": "loaded",
                })

            elif suffix in [".png", ".jpg", ".jpeg", ".webp"]:
                text, parts = extract_image_text(file)
                documents.extend(parts)

                metadata.append({
                    "file": name,
                    "type": "CHART/IMAGE",
                    "rows": None,
                    "status": "loaded",
                })

            else:
                metadata.append({
                    "file": name,
                    "type": suffix,
                    "rows": None,
                    "status": "unsupported",
                })

        except Exception as exc:
            metadata.append({
                "file": name,
                "type": suffix,
                "rows": None,
                "status": f"ERROR: {exc}",
            })

    return documents, tables, metadata


# ============================================================
# CHUNKING
# ============================================================

def chunk_text(text: str, size: int, overlap: int) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()

    if not text:
        return []

    if overlap >= size:
        overlap = max(0, size // 5)

    chunks = []
    start = 0

    while start < len(text):
        end = min(len(text), start + size)

        # Prefer ending at sentence/word boundaries.
        if end < len(text):
            boundary = max(
                text.rfind(". ", start, end),
                text.rfind("; ", start, end),
                text.rfind(" ", start, end),
            )

            if boundary > start + int(size * 0.55):
                end = boundary + 1

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(0, end - overlap)

    return chunks


def build_chunks(documents: List[Dict[str, Any]], size: int, overlap: int):
    all_chunks = []

    for doc in documents:
        pieces = chunk_text(doc["text"], size, overlap)

        for i, piece in enumerate(pieces):
            all_chunks.append({
                "text": piece,
                "source": doc["source"],
                "page": doc.get("page"),
                "type": doc.get("type"),
                "chunk_id": i,
            })

    return all_chunks


# ============================================================
# EMBEDDING / CHROMA VECTOR DATABASE
# ============================================================

CHROMA_DIR = os.getenv("CHROMA_PERSIST_DIR", "./chroma_db")
CHROMA_COLLECTION = "nexamart_rag_documents"


@st.cache_resource
def get_embedder():
    """Load the local SentenceTransformer embedding model once."""
    if SentenceTransformer is None:
        return None

    return SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")


@st.cache_resource
def get_chroma_client():
    """Create a persistent ChromaDB client."""
    if chromadb is None:
        raise RuntimeError(
            "chromadb is not installed. Run: pip install chromadb"
        )

    # Data is persisted on disk in ./chroma_db by default.
    return chromadb.PersistentClient(path=CHROMA_DIR)


def get_chroma_collection(reset: bool = False):
    """Get the persistent RAG collection, optionally recreating it."""
    client = get_chroma_client()

    if reset:
        try:
            client.delete_collection(CHROMA_COLLECTION)
        except Exception:
            pass

    return client.get_or_create_collection(
        name=CHROMA_COLLECTION,
        metadata={"hnsw:space": "cosine"},
    )


def create_chroma_index(chunks: List[Dict[str, Any]]):
    """
    Embed chunks and store them in persistent ChromaDB.

    Each chunk keeps source/page/type metadata so retrieved results
    can still be displayed as citations in the Streamlit UI.
    """
    embedder = get_embedder()

    if embedder is None:
        raise RuntimeError(
            "sentence-transformers is not installed. "
            "Run: pip install sentence-transformers"
        )

    if not chunks:
        return None, 0

    collection = get_chroma_collection(reset=True)

    texts = [c["text"] for c in chunks]

    embeddings = embedder.encode(
        texts,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    ids = []
    metadatas = []

    for i, chunk in enumerate(chunks):
        # Stable ID based on source/page/chunk/text.
        raw_id = (
            f"{chunk.get('source', '')}|"
            f"{chunk.get('page', '')}|"
            f"{chunk.get('chunk_id', i)}|"
            f"{chunk.get('text', '')}"
        )
        ids.append(hashlib.sha256(raw_id.encode("utf-8")).hexdigest())

        # Chroma metadata values should be scalar and not None.
        metadatas.append({
            "source": str(chunk.get("source") or ""),
            "page": str(chunk.get("page") or ""),
            "type": str(chunk.get("type") or ""),
            "chunk_id": str(chunk.get("chunk_id", i)),
        })

    collection.add(
        ids=ids,
        documents=texts,
        embeddings=np.asarray(embeddings).tolist(),
        metadatas=metadatas,
    )

    return collection, len(chunks)


def retrieve(question: str, chunks, embeddings, k: int):
    """
    Retrieve relevant chunks from persistent ChromaDB.

    The old in-memory NumPy cosine search has been replaced by
    Chroma similarity search. The chunks/embeddings arguments are
    retained in the signature so the rest of the app remains compatible.
    """
    if not question.strip():
        return []

    embedder = get_embedder()
    if embedder is None:
        return []

    try:
        collection = get_chroma_collection()
    except Exception:
        return []

    if collection.count() == 0:
        return []

    q_embedding = embedder.encode(
        [question],
        normalize_embeddings=True,
        show_progress_bar=False,
    )[0]

    result = collection.query(
        query_embeddings=[np.asarray(q_embedding).tolist()],
        n_results=min(k, collection.count()),
        include=["documents", "metadatas", "distances"],
    )

    documents = result.get("documents", [[]])[0]
    metadatas = result.get("metadatas", [[]])[0]
    distances = result.get("distances", [[]])[0]

    results = []

    for doc, metadata, distance in zip(documents, metadatas, distances):
        metadata = metadata or {}

        # With cosine distance, similarity is approximately 1 - distance.
        similarity = max(0.0, min(1.0, 1.0 - float(distance)))

        results.append({
            "text": doc,
            "source": metadata.get("source", ""),
            "page": metadata.get("page") or None,
            "type": metadata.get("type", ""),
            "chunk_id": metadata.get("chunk_id", ""),
            "score": similarity,
        })

    return results


# ============================================================
# STRUCTURED CSV ANALYTICS
# ============================================================

def make_table_summary(df: pd.DataFrame, filename: str) -> str:
    lines = [f"DATASET: {filename}", f"Rows: {len(df)}"]

    # Revenue
    revenue_col = find_column(
        df,
        [
            "revenue",
            "sales",
            "sales_amount",
            "net_revenue",
            "net_sales",
            "total_sales",
            "amount",
            "net_amount",
            "revenue_aed",
        ],
    )

    expense_col = find_column(
        df,
        [
            "expense",
            "expenses",
            "total_expense",
            "operating_expense",
            "cost",
            "cogs",
            "cost_of_goods_sold",
        ],
    )

    profit_col = find_column(
        df,
        [
            "profit",
            "gross_profit",
            "net_profit",
            "operating_profit",
        ],
    )

    product_col = find_column(
        df,
        [
            "product",
            "product_name",
            "product_id",
            "sku",
            "item",
        ],
    )

    customer_col = find_column(
        df,
        [
            "customer",
            "customer_name",
            "customer_id",
            "client",
        ],
    )

    inventory_col = find_column(
        df,
        [
            "inventory",
            "stock",
            "stock_quantity",
            "quantity_in_stock",
            "units_in_stock",
            "closing_stock",
        ],
    )

    category_col = find_column(
        df,
        [
            "category",
            "product_category",
            "department",
            "segment",
        ],
    )

    lines.append(
        "Detected columns: "
        f"revenue={revenue_col}, expense={expense_col}, "
        f"profit={profit_col}, product={product_col}, "
        f"customer={customer_col}, inventory={inventory_col}, "
        f"category={category_col}"
    )

    # Overall numerical totals.
    for col, label in [
        (revenue_col, "Total revenue"),
        (expense_col, "Total expenses/cost"),
        (profit_col, "Total profit"),
    ]:
        if col:
            values = pd.to_numeric(df[col], errors="coerce").fillna(0)
            lines.append(f"{label}: {values.sum():,.2f}")

    # Product revenue ranking.
    if product_col and revenue_col:
        temp = df[[product_col, revenue_col]].copy()
        temp[revenue_col] = pd.to_numeric(
            temp[revenue_col], errors="coerce"
        ).fillna(0)

        ranking = (
            temp.groupby(product_col)[revenue_col]
            .sum()
            .sort_values(ascending=False)
            .head(10)
        )

        lines.append("\nTOP PRODUCTS BY REVENUE:")
        for product, value in ranking.items():
            lines.append(f"- {product}: {value:,.2f}")

    # Category performance.
    if category_col and revenue_col:
        temp = df[[category_col, revenue_col]].copy()
        temp[revenue_col] = pd.to_numeric(
            temp[revenue_col], errors="coerce"
        ).fillna(0)

        ranking = (
            temp.groupby(category_col)[revenue_col]
            .sum()
            .sort_values(ascending=False)
            .head(10)
        )

        lines.append("\nCATEGORY REVENUE:")
        for category, value in ranking.items():
            lines.append(f"- {category}: {value:,.2f}")

    # Customer ranking.
    if customer_col and revenue_col:
        temp = df[[customer_col, revenue_col]].copy()
        temp[revenue_col] = pd.to_numeric(
            temp[revenue_col], errors="coerce"
        ).fillna(0)

        ranking = (
            temp.groupby(customer_col)[revenue_col]
            .sum()
            .sort_values(ascending=False)
            .head(10)
        )

        lines.append("\nTOP CUSTOMERS BY REVENUE:")
        for customer, value in ranking.items():
            lines.append(f"- {customer}: {value:,.2f}")

    # Low inventory.
    if product_col and inventory_col:
        temp = df[[product_col, inventory_col]].copy()
        temp[inventory_col] = pd.to_numeric(
            temp[inventory_col], errors="coerce"
        )

        low = (
            temp.dropna()
            .sort_values(inventory_col)
            .head(15)
        )

        lines.append("\nLOWEST INVENTORY PRODUCTS:")
        for _, row in low.iterrows():
            lines.append(
                f"- {row[product_col]}: {row[inventory_col]:,.0f} units"
            )

    return "\n".join(lines)


def build_structured_context(tables: Dict[str, pd.DataFrame]) -> str:
    if not tables:
        return "No CSV datasets have been uploaded."

    sections = []

    for filename, df in tables.items():
        try:
            sections.append(make_table_summary(df, filename))
        except Exception as exc:
            sections.append(
                f"DATASET: {filename}\nSummary generation failed: {exc}"
            )

    return "\n\n".join(sections)


def direct_csv_answer(question: str, tables: Dict[str, pd.DataFrame]) -> str:
    """
    Deterministic calculations for common business questions.
    Returns an empty string when no strong calculation can be inferred.
    """
    q = question.lower()
    answers = []

    years = detect_requested_years(question)

    for filename, df in tables.items():
        revenue_col = find_column(
            df,
            [
                "revenue",
                "sales",
                "sales_amount",
                "net_revenue",
                "net_sales",
                "total_sales",
                "amount",
                "net_amount",
            ],
        )

        expense_col = find_column(
            df,
            [
                "expense",
                "expenses",
                "total_expense",
                "operating_expense",
                "cost",
                "cogs",
            ],
        )

        profit_col = find_column(
            df,
            [
                "profit",
                "gross_profit",
                "net_profit",
                "operating_profit",
            ],
        )

        product_col = find_column(
            df,
            ["product", "product_name", "product_id", "sku", "item"],
        )

        customer_col = find_column(
            df,
            ["customer", "customer_name", "customer_id", "client"],
        )

        inventory_col = find_column(
            df,
            [
                "inventory",
                "stock",
                "stock_quantity",
                "quantity_in_stock",
                "units_in_stock",
                "closing_stock",
            ],
        )

        category_col = find_column(
            df,
            ["category", "product_category", "department", "segment"],
        )

        date_col = None
        for col in detect_date_columns(df):
            date_col = col
            break

        work = df.copy()

        # Parse date column when possible.
        if date_col:
            work["_rag_date"] = pd.to_datetime(
                work[date_col],
                errors="coerce",
            )

        # Handle explicit year questions.
        if years and date_col:
            work = work[
                work["_rag_date"].dt.year.isin(years)
            ].copy()

        # Revenue.
        if revenue_col and (
            "revenue" in q
            or "sales" in q
            or "financial result" in q
        ):
            values = pd.to_numeric(
                work[revenue_col],
                errors="coerce",
            ).fillna(0)

            if "total" in q or "what was" in q or "how much" in q:
                total = values.sum()

                period = (
                    ", ".join(map(str, years))
                    if years
                    else "uploaded dataset period"
                )

                answers.append(
                    f"From `{filename}`, revenue for {period} "
                    f"is **{total:,.2f}**."
                )

        # Expenses.
        if expense_col and (
            "expense" in q or "expenses" in q or "major expense" in q
        ):
            temp = work.copy()
            temp["_value"] = pd.to_numeric(
                temp[expense_col],
                errors="coerce",
            ).fillna(0)

            if category_col:
                grouped = (
                    temp.groupby(category_col)["_value"]
                    .sum()
                    .sort_values(ascending=False)
                    .head(10)
                )
                answers.append(
                    f"Major expense categories in `{filename}`:\n"
                    + "\n".join(
                        f"- {idx}: {value:,.2f}"
                        for idx, value in grouped.items()
                    )
                )
            else:
                answers.append(
                    f"Total expenses/cost in `{filename}`: "
                    f"**{temp['_value'].sum():,.2f}**."
                )

        # Profit.
        if profit_col and "profit" in q:
            values = pd.to_numeric(
                work[profit_col],
                errors="coerce",
            ).fillna(0)

            answers.append(
                f"Profit in `{filename}` for the requested period: "
                f"**{values.sum():,.2f}**."
            )

        # Product ranking.
        if product_col and revenue_col and (
            "which product" in q
            or "top product" in q
            or "most revenue" in q
        ):
            temp = work[[product_col, revenue_col]].copy()
            temp[revenue_col] = pd.to_numeric(
                temp[revenue_col],
                errors="coerce",
            ).fillna(0)

            grouped = (
                temp.groupby(product_col)[revenue_col]
                .sum()
                .sort_values(ascending=False)
            )

            if not grouped.empty:
                top = grouped.head(10)
                answers.append(
                    f"Top products by revenue in `{filename}`:\n"
                    + "\n".join(
                        f"- {idx}: {value:,.2f}"
                        for idx, value in top.items()
                    )
                )

        # Customer ranking.
        if customer_col and revenue_col and (
            "customer" in q and (
                "most" in q
                or "top" in q
                or "purchased" in q
            )
        ):
            temp = work[[customer_col, revenue_col]].copy()
            temp[revenue_col] = pd.to_numeric(
                temp[revenue_col],
                errors="coerce",
            ).fillna(0)

            grouped = (
                temp.groupby(customer_col)[revenue_col]
                .sum()
                .sort_values(ascending=False)
                .head(10)
            )

            answers.append(
                f"Top customers by revenue in `{filename}`:\n"
                + "\n".join(
                    f"- {idx}: {value:,.2f}"
                    for idx, value in grouped.items()
                )
            )

        # Inventory.
        if product_col and inventory_col and (
            "inventory" in q
            or "stock" in q
        ):
            temp = work[[product_col, inventory_col]].copy()
            temp[inventory_col] = pd.to_numeric(
                temp[inventory_col],
                errors="coerce",
            )

            low = (
                temp.dropna()
                .sort_values(inventory_col)
                .head(15)
            )

            answers.append(
                f"Lowest inventory products in `{filename}`:\n"
                + "\n".join(
                    f"- {row[product_col]}: "
                    f"{row[inventory_col]:,.0f} units"
                    for _, row in low.iterrows()
                )
            )

        # Category performance.
        if category_col and revenue_col and (
            "category" in q and (
                "performing" in q
                or "performance" in q
                or "poor" in q
            )
        ):
            temp = work[[category_col, revenue_col]].copy()
            temp[revenue_col] = pd.to_numeric(
                temp[revenue_col],
                errors="coerce",
            ).fillna(0)

            grouped = (
                temp.groupby(category_col)[revenue_col]
                .sum()
                .sort_values()
                .head(10)
            )

            answers.append(
                f"Categories with the lowest revenue in `{filename}`:\n"
                + "\n".join(
                    f"- {idx}: {value:,.2f}"
                    for idx, value in grouped.items()
                )
            )

    return "\n\n".join(dict.fromkeys(answers))


def year_comparison_context(
    question: str,
    tables: Dict[str, pd.DataFrame],
) -> str:
    """
    Compare two explicitly requested years when a date and revenue/profit/
    expense columns are available.
    """
    years = detect_requested_years(question)

    if len(years) < 2:
        return ""

    y1, y2 = years[-2], years[-1]
    results = []

    for filename, df in tables.items():
        date_col = next(iter(detect_date_columns(df)), None)

        revenue_col = find_column(
            df,
            [
                "revenue",
                "sales",
                "sales_amount",
                "net_revenue",
                "net_sales",
                "total_sales",
                "amount",
                "net_amount",
            ],
        )

        profit_col = find_column(
            df,
            ["profit", "gross_profit", "net_profit", "operating_profit"],
        )

        expense_col = find_column(
            df,
            ["expense", "expenses", "total_expense", "operating_expense"],
        )

        if not date_col:
            continue

        if not any([revenue_col, profit_col, expense_col]):
            continue

        temp = df.copy()
        temp["_date"] = pd.to_datetime(temp[date_col], errors="coerce")

        def total_for(col, year):
            if not col:
                return None
            vals = pd.to_numeric(
                temp.loc[temp["_date"].dt.year == year, col],
                errors="coerce",
            ).fillna(0)
            return float(vals.sum())

        metrics = {
            "Revenue": total_for(revenue_col, y1),
            "Profit": total_for(profit_col, y1),
            "Expenses": total_for(expense_col, y1),
        }

        metrics2 = {
            "Revenue": total_for(revenue_col, y2),
            "Profit": total_for(profit_col, y2),
            "Expenses": total_for(expense_col, y2),
        }

        results.append(f"YEAR COMPARISON: {filename}")
        results.append(f"{y1}: {json.dumps(metrics)}")
        results.append(f"{y2}: {json.dumps(metrics2)}")

        for metric in ["Revenue", "Profit", "Expenses"]:
            a = metrics[metric]
            b = metrics2[metric]

            if a is not None and b is not None and a != 0:
                change = ((b - a) / abs(a)) * 100
                results.append(
                    f"{metric} change {y1}->{y2}: "
                    f"{change:.2f}%"
                )

    return "\n".join(results)


# ============================================================
# GROQ
# ============================================================

def get_groq_client(api_key: str):

    if Groq is None:
        raise RuntimeError(
            "groq is not installed. Run: pip install groq"
        )

    key = api_key or os.getenv("GROQ_API_KEY")

    if not key:
        return None

    return Groq(api_key=key)


def answer_with_groq(
    question: str,
    retrieved: List[Dict[str, Any]],
    structured_context: str,
    deterministic_context: str,
    model: str,
    api_key: str,
    temperature: float,
):
    client = get_groq_client(api_key)

    if client is None:
        return (
            "Please enter your Groq API key in the sidebar, or set the "
            "`GROQ_API_KEY` environment variable."
        )

    retrieved_context = "\n\n".join(
        [
            (
                f"SOURCE: {item['source']} | "
                f"PAGE: {item.get('page') or 'N/A'}\n"
                f"{item['text']}"
            )
            for item in retrieved
        ]
    )

    system_prompt = """
You are a Business Intelligence, Data Analyst, and Financial Report Assistant.

Your job is to answer questions using ONLY the uploaded information and
calculated data provided in the context.

Rules:
1. Do not invent financial figures.
2. When a numerical answer is available from the structured CSV analysis,
   prefer that calculated value over a vague statement from the LLM.
3. Distinguish calculated facts from explanations/inferences.
4. If the data does not support a conclusion, say so.
5. For "why" questions, identify evidence-supported drivers such as revenue
   changes, cost changes, product mix, category performance, discounts,
   returns, expenses, or inventory. Do not invent causes.
6. For "this year" and "last year", use dates in the uploaded datasets.
   If the current year is not represented, explicitly say so.
7. For comparisons, provide both values and the percentage/absolute change
   when it can be calculated.
8. For policy questions, quote/paraphrase only what the uploaded policy says.
9. Always cite relevant source files at the end.
10. Keep the answer business-friendly and concise, using tables when useful.
"""

    user_prompt = f"""
QUESTION:
{question}

DETERMINISTIC / PANDAS CALCULATIONS:
{deterministic_context or "No direct calculation was generated."}

STRUCTURED DATASET SUMMARIES:
{structured_context}

RETRIEVED DOCUMENT CONTEXT:
{retrieved_context or "No document chunks were retrieved."}

Provide the best grounded answer.
"""

    response = client.chat.completions.create(
        model=model,
        temperature=temperature,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )

    return response.choices[0].message.content


# ============================================================
# SOURCE DISPLAY
# ============================================================

def show_sources(retrieved):
    if not retrieved:
        return

    st.markdown("### 📚 Sources")

    seen = set()

    for item in retrieved:
        key = (item["source"], item.get("page"))

        if key in seen:
            continue

        seen.add(key)

        page = (
            f" — page {item['page']}"
            if item.get("page")
            else ""
        )

        score = (
            f" — similarity {item['score']:.3f}"
            if "score" in item
            else ""
        )

        st.markdown(
            f'<div class="source-box">📄 '
            f'<b>{item["source"]}</b>{page}{score}</div>',
            unsafe_allow_html=True,
        )


# ============================================================
# UI
# ============================================================

st.title("📊 Business Intelligence + Financial RAG Assistant")

st.markdown(
    """
Upload company documents and datasets, then ask questions about:

**Revenue • Expenses • Profit • Products • Customers • Inventory •
Categories • Financial Reports • Policies • Year-over-Year Performance**
"""
)

uploaded_files = st.file_uploader(
    "📂 ⬆️ Upload Your Files 🌸😊",
    type=[
        "pdf",
        "csv",
        "xlsx",
        "xls",
        "docx",
        "doc",
        "txt",
        "md",
        "png",
        "jpg",
        "jpeg",
        "webp",
    ],
    accept_multiple_files=True,
)

col1, col2 = st.columns([1, 1])

with col1:
    index_button = st.button(
        "🚀 Process & Index Files",
        type="primary",
        use_container_width=True,
    )

with col2:
    clear_button = st.button(
        "🗑️ Clear Knowledge Base",
        use_container_width=True,
    )

if clear_button:
    st.session_state.documents = []
    st.session_state.chunks = []
    st.session_state.embeddings = None
    st.session_state.chroma_collection = None
    st.session_state.tables = {}

    # Remove the persistent Chroma collection as well.
    if chromadb is not None:
        try:
            get_chroma_collection(reset=True)
        except Exception:
            pass
    st.session_state.file_metadata = []
    st.session_state.messages = []
    st.session_state.indexed = False
    st.rerun()


if index_button:
    if not uploaded_files:
        st.warning("Please upload at least one file.")
    else:
        with st.spinner("Reading and indexing your files..."):
            documents, tables, metadata = load_uploaded_files(
                uploaded_files
            )

            chunks = build_chunks(
                documents,
                chunk_size,
                chunk_overlap,
            )

            chroma_collection, vector_count = create_chroma_index(chunks)

            st.session_state.documents = documents
            st.session_state.tables = tables
            st.session_state.chunks = chunks
            st.session_state.embeddings = None
            st.session_state.chroma_collection = chroma_collection
            st.session_state.file_metadata = metadata
            st.session_state.indexed = True

        st.success(
            f"Indexed {len(documents)} source sections into "
            f"{vector_count} chunks and stored them in ChromaDB."
        )
        st.info(f"💾 ChromaDB persistent storage: `{CHROMA_DIR}`")


if st.session_state.indexed:
    st.markdown("---")

    st.subheader("📁 Uploaded Files")

    metadata_df = pd.DataFrame(st.session_state.file_metadata)

    if not metadata_df.empty:
        st.dataframe(
            metadata_df,
            use_container_width=True,
            hide_index=True,
        )

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "Files",
        len(st.session_state.file_metadata),
    )

    c2.metric(
        "CSV Tables",
        len(st.session_state.tables),
    )

    c3.metric(
        "RAG Chunks",
        len(st.session_state.chunks),
    )

    try:
        chroma_count = get_chroma_collection().count()
    except Exception:
        chroma_count = 0

    c4.metric(
        "Chroma Vectors",
        chroma_count,
    )

    with st.expander("🔎 Dataset summaries"):
        for filename, df in st.session_state.tables.items():
            st.markdown(f"#### {filename}")
            st.code(
                make_table_summary(df, filename),
                language="text",
            )


# ============================================================
# CHAT HISTORY
# ============================================================

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

        if message.get("sources"):
            show_sources(message["sources"])


# ============================================================
# CHAT INPUT
# ============================================================

question = st.chat_input(
    "Ask about revenue, profit, products, customers, inventory, policies..."
)

if question:
    if not st.session_state.indexed:
        st.warning("Upload and process your files first.")
        st.stop()

    st.session_state.messages.append({
        "role": "user",
        "content": question,
    })

    with st.chat_message("user"):
        st.markdown(question)

    with st.chat_message("assistant"):
        with st.spinner("Analyzing your data and documents..."):

            retrieved = retrieve(
                question,
                st.session_state.chunks,
                st.session_state.embeddings,
                top_k,
            )

            structured_context = build_structured_context(
                st.session_state.tables
            )

            direct_context = direct_csv_answer(
                question,
                st.session_state.tables,
            )

            comparison_context = year_comparison_context(
                question,
                st.session_state.tables,
            )

            if comparison_context:
                direct_context = (
                    direct_context
                    + "\n\n"
                    + comparison_context
                ).strip()

            try:
                answer = answer_with_groq(
                    question=question,
                    retrieved=retrieved,
                    structured_context=structured_context,
                    deterministic_context=direct_context,
                    model=selected_model,
                    api_key=groq_api_input,
                    temperature=temperature,
                )
            except Exception as exc:
                answer = (
                    "I could not complete the LLM response.\n\n"
                    f"**Error:** `{exc}`\n\n"
                    "Check your Groq API key, selected model, and "
                    "internet/API access."
                )

            st.markdown(answer)

            if retrieved:
                show_sources(retrieved)

            st.session_state.messages.append({
                "role": "assistant",
                "content": answer,
                "sources": retrieved,
            })


# ============================================================
# SAMPLE QUESTIONS
# ============================================================

with st.expander("💡 Example questions"):
    st.markdown(
        """
### Financial
- What was total revenue in 2025?
- What was revenue in last year?
- What was revenue in this year?
- What were the major expenses?
- Summarize the financial results.
- Compare 2024 and 2025.
- Why did profit decrease?

### Sales / Products
- Which product generated the most revenue?
- Which product categories are performing poorly?
- Which products have low inventory?
- Which customers purchased the most?

### Business / Policy
- What does the company policy say about discounts?
- Compare sales performance with the pricing policy.
- Which products had strong sales but low inventory?

### Cross-document
- Compare the annual report with the sales CSV.
- Does the sales data support the financial report?
- Which category contributed most to revenue?
- What business factors appear to explain the change in profit?
"""
    )
